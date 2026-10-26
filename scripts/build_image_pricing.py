#!/usr/bin/env python3
"""Extend the approved image pricing workbook; preserve original ZIP parts byte-for-byte.
Run: python3 scripts/build_image_pricing.py [--skill-dir PATH]
Requires openpyxl and the official xlsx skill (LibreOffice recalc + QA).
"""
import argparse
import copy
import datetime
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / '视频模型成本与售价利润表.xlsx'
BACKUP = ROOT / '视频模型成本与售价利润表_backup.xlsx'
DATE = '2026-10-25'
N = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
Q = '{'+N['s']+'}'
URLQ = 'https://docs.bailian.console.aliyun.com/zh/model-studio/model-pricing'
URLS = 'https://ark.volcengine.com/docs/model-pricing'

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--skill-dir', default=os.environ.get('XLSX_SKILL_DIR', '/Users/fatballfish/.mikiko/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.2.0/skills/xlsx'))
    args = ap.parse_args()
    skill = Path(args.skill_dir)
    sys.path.insert(0, str(skill / 'templates'))
    from base import copy_style, auto_fit_columns, auto_fit_row_heights
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    from base import CF_NEGATIVE_FONT, CF_NEGATIVE_FILL

    if not BACKUP.exists():
        shutil.copy2(BOOK, BACKUP)
    original = load_workbook(BACKUP)
    assert original.sheetnames == ['说明','费率与参数','成本与现售价','建议售价清单','利润率曲线']
    wb = load_workbook(BACKUP)
    print('Actual parameters:', [wb['费率与参数'][f'B{i}'].value for i in range(2,7)], flush=True)
    wb.properties.creator = 'Z.ai'
    headers_style = wb['成本与现售价']['A3']
    body_style = wb['成本与现售价']['A4']
    input_style = wb['费率与参数']['B2']
    p = "'费率与参数'!"
    new_sheets = []
    rate_rows = {}

    with tempfile.TemporaryDirectory(prefix='image-pricing-') as td:
        td = Path(td)
        draft = td / 'pricing.xlsx'
        def qa():
            wb.save(draft)
            for cmd in ('recalc','audit','scan'):
                r = subprocess.run([sys.executable, str(skill/'xlsx.py'),cmd,str(draft)], capture_output=True,text=True)
                result = json.loads(r.stdout)
                print(cmd, json.dumps({k:v for k,v in result.items() if not isinstance(v,list)},ensure_ascii=False), flush=True)
                if cmd != 'scan': assert r.returncode == 0, (cmd,r.stdout,r.stderr)
                if cmd == 'recalc':
                    result = json.loads(r.stdout)
                    assert result.get('status') == 'success', result
            cached = load_workbook(draft,data_only=True)
            for name in new_sheets:
                for row in wb[name]:
                    for c in row:
                        if c.data_type == 'f':
                            v = cached[name][c.coordinate]
                            assert v.data_type != 'e', (name,c.coordinate,v.value)
                            assert v.value is not None, (name,c.coordinate,'missing cache')
            return cached

        def sheet(name, title, headers):
            ws=wb.create_sheet(name); new_sheets.append(name)
            ws['B2']=title; copy_style(wb['说明']['A2'],ws['B2'])
            ws.merge_cells(start_row=2,start_column=2,end_row=2,end_column=len(headers)+1)
            for j,h in enumerate(headers,2):
                c=ws.cell(4,j,h); copy_style(headers_style,c)
                c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            ws.freeze_panes='D5'; ws.sheet_view.showGridLines=False
            ws.column_dimensions['A'].width=3
            ws.row_dimensions[2].height=32
            return ws

        def finish(ws):
            for row in ws.iter_rows(min_row=5):
                for c in row[1:]:
                    copy_style(body_style,c)
                    c.alignment=Alignment(horizontal='right' if c.data_type=='f' or isinstance(c.value,(int,float)) else 'left',vertical='center',wrap_text=True)
                    if isinstance(c.value,(int,float)):
                        copy_style(input_style,c)
                        c.comment=Comment('来源：专项PRD/design及本表来源列；可调场景数量为分析假设。核验日期 '+DATE,'Z.ai')
                    if c.data_type=='f' or isinstance(c.value,float): c.number_format='0.0000'
            auto_fit_columns(ws,min_width=10,max_width=38,header_row=4,data_start_row=5)
            auto_fit_row_heights(ws,header_row=4,data_start_row=5)
            ws.auto_filter.ref=f'B4:{ws.cell(ws.max_row,ws.max_column).coordinate}'
            ws.print_title_rows='1:4'
            ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
            qa()

        rates=sheet('图片费率','图片官方牌价｜CNY｜未知值禁止按零计价',[
            '费率键','模型','计费项','档位/条件','费率（CNY/单位）','单位','状态','来源链接','核验日期','口径'])
        def rate(key,model,item,tier,value,unit,note,url=URLQ):
            r=rates.max_row+1
            rates.append([None,key,model,item,tier,value,unit,'待核实' if value is None else '已核验',url,DATE,note])
            rate_rows[key]=r
            rates.cell(r,9).hyperlink=url
        models=['qwen-image-3.0','qwen-image-2.1-pro','qwen-image-2.1-turbo','qwen-image-3.0-pro','doubao-seedream-5-0-pro-260628','doubao-seedream-5-0-flash-260915']
        for model,output,inp in [(models[0],.18,.02),(models[1],.25,0),(models[2],.10,0),(models[3],.25,.02)]:
            rate(model+'|input',model,'input','全部',inp,'张','2.1输入未标价，按官方未注明输入价则输入免费通则；3.0输入按实际传入张数。')
            for tier in ('1K','2K'):
                rate(model+'|'+tier,model,'output',tier,.50 if model==models[3] and tier=='2K' else output,'张','3.0按输出面积≤2250000像素为1K；2.1无分档价，两尺寸同价。3.0-pro为可选预设。')
        pro,flash=models[4:]
        rate(pro+'|input',pro,'input','第2张起',.02,'张','首输入图免费；MAX(输入张数-1,0)。',URLS)
        for tier,out,layer in [('1K',.30,.15),('2K',.60,.30)]:
            rate(pro+'|'+tier,pro,'output',tier,out,'张','实际像素≤2610000用低档，否则高档；1K/2K仅本表示例尺寸。',URLS)
            rate(pro+'|layer|'+tier,pro,'layer',tier,layer,'图层','每层按实际像素档位计费，不能按整图尺寸统一计价。',URLS)
            rate(pro+'|base|'+tier,pro,'base',tier,None,'张','底图是否收费、如何归档未明确，完整拆层报价禁用。',URLS)
        rate(flash+'|input',flash,'input','全部',0,'张','官方明确免费，零值为已确认费率。',URLS)
        for tier in ('1K','2K'):
            rate(flash+'|'+tier,flash,'output',tier,.12,'张','普通输出统一单价。',URLS)
            rate(flash+'|layer|'+tier,flash,'layer',tier,None,'图层','Flash拆层费率未公布；不得套用普通输出价。',URLS)
            rate(flash+'|base|'+tier,flash,'base',tier,None,'张','拆层底图收费未明确。',URLS)
        finish(rates)
        def rr(key): return f"'图片费率'!$F${rate_rows[key]}"

        s=sheet('图片场景成本','场景成本｜普通输出1张；拆层1底图+图层；未知完整成本留待核实',[
            '场景ID','模型','操作','尺寸示例','输入（张）','输出（张）','低档层（层）','高档层（层）','输入费（CNY）','输出费（CNY）','已知图层费（CNY）','底图费（CNY）','完整成本（CNY/任务）','状态','说明'])
        scenarios=[]
        for mi,model in enumerate(models):
            for tier in ('1K','2K'):
                for count in (0,1,3 if mi in (0,3) else 10):
                    scenarios.append((model,'generate' if count==0 else 'edit',tier,count,1,0,0))
        for model in (pro,flash):
            for tier in ('1K','2K'):
                for low,high in ((1,0),(3,1),(16,0),(0,16)):
                    scenarios.append((model,'layer_decomposition',tier,1,0,low,high))
        for model,op,tier,count,out,low,high in scenarios:
            r=s.max_row+1; key=f'IMG-{r-4:03d}'
            s.append([None,key,model,op,tier,count,out,low,high])
            qty=f'MAX(F{r}-1,0)' if model==pro else f'F{r}'
            s.cell(r,10,f'={qty}*{rr(model+"|input")}')
            s.cell(r,11,f'=G{r}*{rr(model+"|"+tier)}')
            if op=='layer_decomposition':
                s.cell(r,12,f'=H{r}*{rr(pro+"|layer|1K")}+I{r}*{rr(pro+"|layer|2K")}' if model==pro else '待核实')
                s.cell(r,13,'待核实')
                s.cell(r,14,f'=IF(AND(ISNUMBER(J{r}),ISNUMBER(K{r}),ISNUMBER(L{r}),ISNUMBER(M{r})),SUM(J{r}:M{r}),"待核实")')
                s.cell(r,16,'仅允许1张拆层输入；已知层费为小计，不是完整成本；最大17产物按实际逐层像素结算。')
            else:
                s.cell(r,12,0); s.cell(r,13,0); s.cell(r,14,f'=SUM(J{r}:M{r})')
                s.cell(r,16,'1K示例1024×1024；2K示例2048×2048。单输出避免多输出输入费用口径未明。')
            s.cell(r,15,f'=IF(ISNUMBER(N{r}),"可测算","待核实：禁用建议")')
        for col in ('F','G','H','I'):
            v=DataValidation(type='whole',operator='between',formula1=0,formula2=16); v.errorTitle='数量无效';v.error='请输入0到16的整数并遵守型号及操作限制';v.showErrorMessage=True;s.add_data_validation(v);v.add(f'{col}5:{col}{s.max_row}')
        finish(s)

        price=sheet('图片成本售价','图片建议售价｜仅测算，非线上现行价｜目标毛利按成本对数插值',[
            '场景ID','模型','操作','档位','成本（CNY/任务）','目标毛利率','保本积分','保本收入（CNY）','建议积分/任务','建议收入（CNY）','毛利（CNY）','实际毛利率','建议状态'])
        end=s.max_row
        for r in range(5,end+1):
            for c,src in ((2,'B'),(3,'C'),(4,'D'),(5,'E'),(6,'N')): price.cell(r,c,f"='图片场景成本'!{src}{r}")
            # All unknown costs are text, ignored by MIN/MAX; verified normal costs > 0.
            lo=f"MIN('图片场景成本'!$N$5:$N${end})"; hi=f"MAX('图片场景成本'!$N$5:$N${end})"
            price.cell(r,7,f'=IF(ISNUMBER(F{r}),IFERROR(MIN(0.95,MAX({p}$B$6,{p}$B$5-({p}$B$5-{p}$B$6)*(LN(F{r})-LN({lo}))/(LN({hi})-LN({lo})))),{p}$B$5),"待核实")')
            valid=f'AND(ISNUMBER(F{r}),{p}$B$2>0,{p}$B$3>0,{p}$B$4>=0,{p}$B$6>=0,{p}$B$5<1,{p}$B$5>={p}$B$6)'
            price.cell(r,8,f'=IF({valid},MAX({p}$B$4,CEILING(F{r}/{p}$B$2,{p}$B$3)),"待核实/参数无效")')
            price.cell(r,9,f'=IF(ISNUMBER(H{r}),H{r}*{p}$B$2,"待核实/参数无效")')
            price.cell(r,10,f'=IF({valid},MAX({p}$B$4,CEILING(F{r}/(1-G{r})/{p}$B$2,{p}$B$3)),"待核实/参数无效")')
            price.cell(r,11,f'=IF(ISNUMBER(J{r}),J{r}*{p}$B$2,"待核实/参数无效")')
            price.cell(r,12,f'=IF(ISNUMBER(K{r}),K{r}-F{r},"待核实/参数无效")')
            price.cell(r,13,f'=IF(ISNUMBER(K{r}),IFERROR(L{r}/K{r},0),"待核实/参数无效")')
            price.cell(r,14,f'=IF(ISNUMBER(J{r}),"建议候选，需后台审核","禁用")')
        finish(price)
        for r in range(5,end+1):
            for c in ('G','M'): price[f'{c}{r}'].number_format='0.00%'

        sens=sheet('图片兑换敏感性','兑换敏感性｜固定基准建议积分，兑换率改变时重算收入毛利；同时列重新定价积分',[
            '场景ID','基准建议积分','成本（CNY）','兑换倍率（假设）','兑换价（CNY/积分）','固定积分收入（CNY）','固定积分毛利（CNY）','固定积分毛利率','维持目标毛利需积分','重新定价收入（CNY）','状态'])
        for sr in range(5,end+1):
            for mult in (.5,1,2):
                r=sens.max_row+1
                sens.append([None,f"='图片成本售价'!B{sr}",f"='图片成本售价'!J{sr}",f"='图片成本售价'!F{sr}",mult,f'={p}$B$2*E{r}'])
                sens.cell(r,7,f'=IF(ISNUMBER(C{r}),C{r}*F{r},"待核实")')
                sens.cell(r,8,f'=IF(ISNUMBER(G{r}),G{r}-D{r},"待核实")')
                sens.cell(r,9,f'=IF(ISNUMBER(G{r}),IFERROR(H{r}/G{r},0),"待核实")')
                sens.cell(r,10,f'=IF(AND(ISNUMBER(C{r}),F{r}>0,{p}$B$3>0),MAX({p}$B$4,CEILING(D{r}/(1-\'图片成本售价\'!G{sr})/F{r},{p}$B$3)),"待核实/参数无效")')
                sens.cell(r,11,f'=IF(ISNUMBER(J{r}),J{r}*F{r},"待核实")')
                sens.cell(r,12,f'=IF(ISNUMBER(C{r}),"固定积分存在兑换风险","禁用")')
        finish(sens)
        for r in range(5,sens.max_row+1): sens[f'I{r}'].number_format='0.00%'
        sens.conditional_formatting.add(f'H5:H{sens.max_row}',CellIsRule(operator='lessThan',formula=['0'],font=CF_NEGATIVE_FONT,fill=CF_NEGATIVE_FILL))

        cfg=sheet('图片后台建议配置','后台建议配置｜仅建议，不写入production；费率单位CNY，售价单位积分',[
            '模型','预设性质','操作','输入上限（张）','普通输出（张）','产物上限','输入费率（CNY/张）','1K输出（CNY/张）','2K输出（CNY/张）','低档层（CNY/层）','高档层（CNY/层）','底图费','建议启用','结算/能力说明'])
        for mi,model in enumerate(models):
            for op in ('generate/edit','layer_decomposition') if mi>=4 else ('generate/edit',):
                r=cfg.max_row+1
                cfg.append([None,model,'额外可选' if mi==3 else '正式接入候选',op,1 if op=='layer_decomposition' else 3 if mi in (0,3) else 10,1 if op!='layer_decomposition' else 0,17 if op=='layer_decomposition' else 1,f'={rr(model+"|input")}',f'={rr(model+"|1K")}',f'={rr(model+"|2K")}'])
                for c,tier in ((11,'1K'),(12,'2K')):
                    cfg.cell(r,c,f'={rr(pro+"|layer|"+tier)}' if model==pro else '待核实' if model==flash else '不适用')
                cfg.cell(r,13,'待核实' if mi>=4 else '不适用')
                cfg.cell(r,14,'禁用：完整拆层费率待核实' if op=='layer_decomposition' else '可选，默认不启用' if mi==3 else '审核后可启用普通生成/编辑')
                cfg.cell(r,15,'冻结费率、售价和积分兑换率；失败释放，按实际结算且幂等。Pro首输入免费；拆层按各层像素，禁止用普通count代替产物。普通多输出暂不建议；Endpoint ID需后台真实配置。')
        finish(cfg)

        rev=sheet('图片校验','图片计算交叉校验｜PASS方可交付', ['校验项','期望','实际','结果'])
        checks=[('兑换参数链接',f'={p}$B$2',f"='图片兑换敏感性'!F6"),('最低积分',f'={p}$B$4',f'={p}$B$4'),('最低成本',.10,"=MIN('图片场景成本'!N5:N40)"),('Pro三层低档一层高档小计',.75,"='图片场景成本'!L42"),('未知完整拆层数量',16,f'=COUNTIF(\'图片场景成本\'!N5:N{end},"待核实")'),('待核实费率非零占位',6,f'=COUNTBLANK(\'图片费率\'!F5:F{rates.max_row})'),('场景数量',len(scenarios),f"=COUNTA('图片场景成本'!B5:B{end})")]
        for name,expected,actual in checks:
            r=rev.max_row+1;rev.append([None,name,expected,actual,f'=IF(C{r}=D{r},"PASS","FAIL")'])
        finish(rev)
        caches=qa()
        assert all(caches['图片校验'].cell(r,5).value=='PASS' for r in range(5,rev.max_row+1)),[(r,caches['图片校验'].cell(r,5).value) for r in range(5,rev.max_row+1)]
        # Use pristine openpyxl output for formulas/styles, transplant only LibreOffice caches.
        pristine=td/'pristine.xlsx';wb.save(pristine)
        with zipfile.ZipFile(BACKUP) as z: orig={n:z.read(n) for n in z.namelist()}
        with zipfile.ZipFile(pristine) as z: final={n:z.read(n) for n in z.namelist()}
        with zipfile.ZipFile(draft) as z:
            for idx,name in enumerate(wb.sheetnames,1):
                if name not in new_sheets: continue
                part=f'xl/worksheets/sheet{idx}.xml'
                tree=ET.fromstring(final[part]); calc=ET.fromstring(z.read(part))
                lookup={c.attrib['r']:c for c in calc.findall('.//s:c',N)}
                for cell in tree.findall('.//s:c',N):
                    if cell.find(Q+'f') is None: continue
                    result=lookup[cell.attrib['r']]
                    for child in list(cell):
                        if child.tag in (Q+'v',Q+'is'):cell.remove(child)
                    t=result.attrib.get('t')
                    if t=='inlineStr':
                        cell.set('t','str');v=ET.SubElement(cell,Q+'v');v.text=''.join(result.itertext())
                    else:
                        if t:cell.set('t',t)
                        else:cell.attrib.pop('t',None)
                        v=result.find(Q+'v')
                        if v is not None:cell.append(copy.deepcopy(v))
                final[part]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
        mutable={'xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml','xl/styles.xml','docProps/core.xml'}
        for name,data in orig.items():
            if name not in mutable: final[name]=data
        # openpyxl may normalize original font attributes; restore their exact XML entries.
        old_styles=ET.fromstring(orig['xl/styles.xml'])
        new_styles=ET.fromstring(final['xl/styles.xml'])
        old_num=old_styles.find(Q+'numFmts'); new_num=new_styles.find(Q+'numFmts')
        formats={e.get('formatCode'):e.get('numFmtId') for e in old_num}
        mapping={}
        merged=copy.deepcopy(old_num)
        next_id=max(179,*[int(e.get('numFmtId'))+1 for e in old_num])
        for e in new_num:
            code=e.get('formatCode'); target=formats.get(code)
            if target is None:
                target=str(next_id); next_id+=1; extra=copy.deepcopy(e);extra.set('numFmtId',target);merged.append(extra)
            mapping[e.get('numFmtId')]=target
        merged.set('count',str(len(merged)))
        new_styles.remove(new_num);new_styles.insert(0,merged)
        for xf in new_styles.iter(Q+'xf'):
            if xf.get('numFmtId') in mapping:xf.set('numFmtId',mapping[xf.get('numFmtId')])
        for tag in ('fonts','fills','borders','cellStyleXfs','cellXfs','cellStyles','dxfs'):
            a=old_styles.find(Q+tag); b=new_styles.find(Q+tag)
            if a is None: continue
            assert b is not None and len(b)>=len(a),tag
            extras=list(b)[len(a):]
            for child in list(b): b.remove(child)
            for child in a: b.append(copy.deepcopy(child))
            for child in extras: b.append(child)
            if 'count' in b.attrib: b.set('count',str(len(b)))
        final['xl/styles.xml']=ET.tostring(new_styles,encoding='utf-8',xml_declaration=True)
        # Restore declarations for retained parts (notably original sharedStrings).
        relns='{http://schemas.openxmlformats.org/package/2006/relationships}'
        old_rel=ET.fromstring(orig['xl/_rels/workbook.xml.rels'])
        new_rel=ET.fromstring(final['xl/_rels/workbook.xml.rels'])
        for rel in old_rel:
            if rel.get('Type','').endswith('/sharedStrings') and not any(x.get('Type')==rel.get('Type') for x in new_rel):
                extra=copy.deepcopy(rel);extra.set('Id','rIdPreservedSharedStrings');new_rel.append(extra)
        final['xl/_rels/workbook.xml.rels']=ET.tostring(new_rel,encoding='utf-8',xml_declaration=True)
        old_ct=ET.fromstring(orig['[Content_Types].xml']);new_ct=ET.fromstring(final['[Content_Types].xml'])
        for item in old_ct:
            key='PartName' if 'PartName' in item.attrib else 'Extension'
            if not any(x.get(key)==item.get(key) for x in new_ct):new_ct.append(copy.deepcopy(item))
        final['[Content_Types].xml']=ET.tostring(new_ct,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(BOOK,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in final.items():z.writestr(name,data)
        with zipfile.ZipFile(BOOK) as z:
            assert all(z.read(n)==v for n,v in orig.items() if n not in mutable)
        output=load_workbook(BOOK); data=load_workbook(BOOK,data_only=True)
        for source in original:
            target=output[source.title]
            assert source.max_row==target.max_row and source.max_column==target.max_column
            for row in source:
                for c in row:
                    t=target[c.coordinate]
                    assert c.value==t.value and c._style==t._style,(source.title,c.coordinate)
        for name in new_sheets:
            for row in output[name]:
                for c in row:
                    if c.data_type=='f':assert data[name][c.coordinate].value is not None
        # audit and chart-verify can silently recalc: run release checks on isolated copies.
        release=td/'release-check.xlsx';shutil.copy2(BOOK,release)
        for cmd in ('audit','scan','chart-verify','validate'):
            command=[sys.executable,str(skill/'xlsx.py'),cmd,str(release)]
            if cmd=='chart-verify':command.append('--no-auto-recalc')
            r = subprocess.run(command,capture_output=True,text=True)
            print(cmd,r.stdout,flush=True)
            if cmd == 'chart-verify' and r.returncode:
                baseline_copy=td/'baseline-check.xlsx';shutil.copy2(BACKUP,baseline_copy)
                baseline=subprocess.run([sys.executable,str(skill/'xlsx.py'),cmd,str(baseline_copy),'--no-auto-recalc'],capture_output=True,text=True)
                assert json.loads(r.stdout)==json.loads(baseline.stdout),'New chart regression'
                print('Chart verifier matches unchanged original scatter-chart baseline',flush=True)
            elif cmd != 'scan': assert r.returncode == 0,(cmd,r.stderr)
        print('PASS: original ZIP parts, all original cell values/formulas/styles, cached new formulas and cross-checks. New scenarios:',len(scenarios),flush=True)

if __name__=='__main__':main()
